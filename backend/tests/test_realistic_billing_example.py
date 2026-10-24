"""Replay the varied fictional exports through the actual authorized report API."""
import io
import json
import secrets
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from backend.jobs import work_once
from backend.tests.test_workbench import upload_confirm,approve,accept

BASE=Path('public/examples/realistic-billing')

def test_varied_staffing_output_is_reconciled_and_exported(environment,clients):
    oracle=json.loads((BASE/'expected-results.json').read_text())
    client=clients()
    response=client.post('/api/v1/import-workbench/sessions',json={
        'facility_id':'ALL','service_from':'2026-10-01','service_to':'2026-10-15',
        'source_as_of':'2026-10-21T17:00:00+00:00','billing_cutoff':'2026-10-20T12:00:00+00:00',
        'invoice_complete':True,'work_complete':True,'expected_work_hours':'136',
        'expected_invoice_total':'8328.25','synthetic':True})
    assert response.status_code==200,response.text
    session=response.json()['id'];path='/api/v1/import-workbench/sessions/'+session
    for kind in ['rates','work','invoices']:upload_confirm(client,session,kind,(BASE/(kind+'.csv')).read_bytes())
    response=client.post(path+'/validate',json={'expected_version':client.get(path).json()['version']})
    assert response.status_code==200 and response.json()['valid'],response.text
    approve(clients('DEMO-APPROVER'),session)
    response=accept(client,session);assert response.status_code==200,response.text
    client.headers['Idempotency-Key']=secrets.token_hex(16)
    response=client.post('/api/v1/runs',json={'import_id':response.json()['id']})
    assert response.status_code==200,response.text
    assert work_once(environment)
    state=client.get('/api/v1/workspace').json()['state'];run=state['run']
    assert run['coverage']['received_records']==oracle['work_records']
    assert run['coverage']['evaluated_records']==oracle['evaluated_records']
    assert Decimal(run['coverage']['received_minutes'])/60==Decimal(oracle['work_hours'])
    assert Decimal(run['coverage']['evaluated_minutes'])/60==Decimal(oracle['evaluated_hours'])
    assert run['totals']['expected']==oracle['expected_evaluated']
    assert run['totals']['underbilling']==oracle['underbilling']
    assert run['totals']['overbilling']==oracle['overbilling']
    assert run['totals']['invoiced']==oracle['posted_invoice_total']
    assert len(run['findings'])==oracle['findings'] and len(run['tasks'])==oracle['coverage_gaps']
    assert sum(r['category']=='evaluated_aligned' for r in run['records'])==oracle['aligned_records']
    records={r['work']['id']:r for r in run['records']}
    for identifier,amounts in oracle['exceptions'].items():
        assert [records[identifier][field] for field in ['expected_amount','invoiced_amount','discrepancy_amount']]==amounts
    assert records['RIV-005']['expected_amount'] is None and records['RIV-006']['expected_amount'] is None
    assert len(records['RIV-004']['allocations'])==3
    assert sum(Decimal(a['amount']) for a in records['RIV-004']['allocations'])==Decimal('966.00')
    assert len({r['work']['worked_minutes'] for r in run['records']})>=8
    # Fictional decisions are recorded through the same reviewer route as the UI.
    for finding in run['findings']:
        references=set(finding['shift_ids'])
        action='confirm' if references in [{'BAY-001','BAY-002'},{'BAY-006'},{'RIV-001'}] else 'request_evidence' if references=={'RIV-004'} else None
        if action is None:continue
        client.headers['Idempotency-Key']=secrets.token_hex(16)
        response=client.post('/api/v1/findings/'+finding['id']+'/decisions',json={
            'action':action,'reason':'Fictional training review: inspect the dated terms and posted statement before acting.',
            'expected_fingerprint':finding['fingerprint'],'expected_version':1,'run_id':run['id'],
            'owner':'DEMO-REVIEWER' if action=='request_evidence' else None})
        assert response.status_code==200,(finding['id'],action,response.text)
    target=Path('artifacts/reports');target.mkdir(parents=True,exist_ok=True)
    response=client.get('/api/v1/runs/'+run['id']+'/report.xlsx')
    assert response.status_code==200,response.text
    book=load_workbook(io.BytesIO(response.content));overview=dict(list(book['Overview'].values)[1:])
    assert overview['Net invoiced on evaluated work (USD)']==Decimal(oracle['evaluated_invoice_total'])
    assert overview['Reviewer-confirmed underbilling difference (USD)']==Decimal(oracle['confirmed_underbilling'])
    assert overview['Reviewer-confirmed overbilling difference (USD)']==Decimal(oracle['confirmed_overbilling'])
    labels=[row[2] for row in list(book['Findings'].values)[1:]]
    assert labels.count('Confirmed discrepancy')==3 and labels.count('Evidence requested')==1 and labels.count('Needs review')==3
    assert book['Work calculations'].max_row==19 and book['Source lineage'].max_row==46
    assert all(cell.data_type!='f' for sheet in book for row in sheet for cell in row)
    worker_ids={r['work']['worker_id'] for r in run['records']}
    values='\n'.join(str(c.value) for sheet in book for row in sheet for c in row)
    assert not any(identifier in values for identifier in worker_ids)
    assert 'Fictional training review:' not in values
    (target/'realistic-billing-review.xlsx').write_bytes(response.content)
    response=client.get('/api/v1/runs/'+run['id']+'/report.pdf')
    assert response.status_code==200 and response.content.startswith(b'%PDF-'),response.text
    assert len(__import__('re').findall(rb'/Type /Page\b',response.content))<=7
    (target/'realistic-billing-review.pdf').write_bytes(response.content)
    (target/'realistic-billing-results.json').write_text(json.dumps({'totals':run['totals'],'coverage':run['coverage'],'findings':[{k:f[k] for k in ['id','facility_id','shift_ids','expected_amount','invoiced_amount','discrepancy_amount']} for f in run['findings']]},indent=2))
