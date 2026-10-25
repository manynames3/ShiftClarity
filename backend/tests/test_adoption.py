"""Realistic export mistakes must be visible before financial review."""
import csv,io
import pytest
from backend.tabular import inspect_file,TableError
from backend.table_mapping import normalize_table

def table(headers,rows,category='work'):
 s=io.StringIO();w=csv.writer(s);w.writerow(headers);w.writerows(rows)
 return inspect_file(s.getvalue().encode(),'fictional.csv',category)

def config(t,unit='hours'):
 return {'worksheet':t['worksheet'],'header_row':t['header_row'],'mapping':t['mapping'],'constants':{},'value_maps':t['value_maps'],'date_format':'ISO','decimal_format':'dot','time_unit':unit,'end_date_mode':'exclusive','confirmed':True}

def test_minute_header_cannot_silently_be_confirmed_as_hours():
 t=table(['Shift ID','Facility','Employee ID','Role','Pay Type','Shift Date','Minutes','Approval Status','Approved Date'],[['W1','FAC-A','E1','RN','Regular','2026-10-05','450','Yes','2026-10-06']])
 with pytest.raises(TableError,match='minutes') as error:normalize_table(t,'work',config(t),'work.csv','sha')
 assert error.value.code=='TIME_UNIT_MISMATCH'
 result=normalize_table(t,'work',config(t,'minutes'),'work.csv','sha')
 assert result['records'][0]['worked_minutes']==450 and not result['errors']

def test_worker_suggestions_require_reference_style_headers():
 t=table(['Shift ID','Employee','Candidate ID','Bill Transaction Date','Bill Qty','Bill Earn Code','Bill Transaction Status','Job Title'],[['W1','Fictional','101','2026-10-05','8','Reg','Approved','RN']])
 assert t['mapping']['worker_id']=='Candidate ID'
 assert t['mapping']['hours']=='Bill Qty'
 assert t['mapping']['service_date']=='Bill Transaction Date'
 assert t['mapping']['component']=='Bill Earn Code'
 t=table(['Shift ID','Employee','Role'],[['W1','Fictional','RN']])
 assert 'worker_id' not in t['mapping']

def test_sensitive_header_diagnostic_identifies_columns_without_cell_values():
 with pytest.raises(TableError) as error:table(['Shift ID','Candidate Name','Phone Number'],[['W1','PRIVATE-CANARY-VALUE','PRIVATE-NUMBER']])
 assert error.value.code=='SENSITIVE_FIELDS'
 assert 'Candidate Name' in str(error.value) and 'Phone Number' in str(error.value)
 assert 'PRIVATE' not in str(error.value)


def test_non_hourly_amount_remains_visible_without_invented_hours():
 t=table(['Line ID','Invoice ID','Facility','Amount','Status','Type','Billing basis'],[['F1','INV1','FAC-A','75.00','Posted','Invoice','Non-hourly']], 'invoices')
 result=normalize_table(t,'invoices',config(t),'invoice.csv','sha')
 assert result['errors']==[]
 line=result['records'][0]
 assert line['billing_basis']=='non_hourly' and line['amount']=='75.00'
 assert line['minutes']==0 and line['unit_rate']=='0.0000' and line['shift_ids']==[]

def test_observed_approved_status_does_not_invent_an_approval_event_date():
 t=table(['Shift ID','Facility','Employee ID','Role','Pay Type','Shift Date','Hours','Approval Status'],[['W1','FAC-A','E1','RN','Regular','2026-10-05','8','Yes']])
 cfg=config(t);cfg['constants']['approval_observed_at']='2026-10-12'
 result=normalize_table(t,'work',cfg,'work.csv','sha')
 assert result['errors']==[]
 assert result['records'][0]['approved_at'] is None
 assert result['records'][0]['approval_observed_at']=='2026-10-12T00:00:00+00:00'


def test_agency_exports_survive_the_real_private_import_approval_and_report_pipeline(environment,clients):
    import io,secrets
    from pathlib import Path
    from decimal import Decimal
    from openpyxl import load_workbook
    from backend.jobs import work_once
    from backend.tests.test_workbench import approve,accept
    client=clients()
    response=client.post('/api/v1/import-workbench/sessions',json={'facility_id':'FAC-ALPHA','service_from':'2026-10-05','service_to':'2026-10-11','source_as_of':'2026-10-14T00:00:00Z','billing_cutoff':'2026-10-13T00:00:00Z','invoice_complete':True,'work_complete':True,'expected_work_hours':'69','expected_invoice_total':'3501.00','synthetic':True})
    assert response.status_code==200,response.text
    session=response.json()['id'];path='/api/v1/import-workbench/sessions/'+session
    for kind in ['rates','work','invoices']:
        data=Path('public/examples/agency-evaluation/'+kind+'.csv').read_bytes()
        response=client.post(path+'/files',params={'category':kind,'filename':kind+'.csv'},content=data,headers={'Content-Type':'application/octet-stream','X-Data-Classification':'synthetic'})
        assert response.status_code==200,response.text
        file_id=response.json()['id'];t=inspect_file(data,kind+'.csv',kind);cfg=config(t,'minutes')
        response=client.post(path+'/files/'+file_id+'/mapping',json={**cfg,'expected_version':client.get(path).json()['version']})
        assert response.status_code==200 and not response.json()['errors'],response.text
    response=client.post(path+'/validate',json={'expected_version':client.get(path).json()['version']})
    assert response.status_code==200 and response.json()['valid'],response.text
    summary=client.get(path).json()['summary'];assert summary['non_hourly_lines']==1 and summary['non_hourly_posted_total']=='25.00'
    approve(clients('DEMO-APPROVER'),session)
    response=accept(client,session);assert response.status_code==200,response.text
    client.headers['Idempotency-Key']=secrets.token_hex(16)
    response=client.post('/api/v1/runs',json={'import_id':response.json()['id']});assert response.status_code==200,response.text
    assert work_once(environment)
    state=client.get('/api/v1/workspace').json()['state'];run=state['run']
    assert run['coverage']['evaluated_records']==6 and len(run['tasks'])==2
    assert run['totals']=={'expected':'3912.00','invoiced':'3501.00','underbilling':'460.00','overbilling':'24.00'}
    response=client.get('/api/v1/runs/'+run['id']+'/report.xlsx');assert response.status_code==200,response.text
    book=load_workbook(io.BytesIO(response.content));assert book['Non-hourly invoice lines']['E2'].value==Decimal('25.00')
    headers=[cell.value for cell in book['Work calculations'][1]];assert headers[-2:]==['Work approval event date','Approved status observed at']
    assert book['Work calculations']['P2'].value=='Unknown'
    target=Path('artifacts/reports');target.mkdir(parents=True,exist_ok=True)
    (target/'agency-evaluation.xlsx').write_bytes(response.content)
    response=client.get('/api/v1/runs/'+run['id']+'/report.pdf');assert response.status_code==200 and response.content.startswith(b'%PDF-'),response.text
    (target/'agency-evaluation.pdf').write_bytes(response.content)


def test_public_example_downloads_are_fictional_and_do_not_expose_private_paths(clients):
    client=clients()
    response=client.get('/examples/agency-evaluation/work.csv')
    assert response.status_code==200 and 'FAC-ALPHA' in response.text
    assert 'Candidate ID' in response.text and 'Candidate Name' not in response.text
    assert client.get('/examples/.local/aws-accounts.json').status_code==404


def test_numeric_excel_cells_ignore_text_decimal_conventions():
    from openpyxl import Workbook
    book=Workbook();sheet=book.active
    sheet.append(['Rate ID','Contract ID','Facility','Role','Rate Type','Bill Rate','Effective From','Authorized On'])
    sheet.append(['R1','C1','FAC-A','RN','Regular',12.125,'2026-10-01','2026-09-30'])
    stream=io.BytesIO();book.save(stream)
    t=inspect_file(stream.getvalue(),'rates.xlsx','rates');cfg=config(t);cfg['decimal_format']='comma'
    result=normalize_table(t,'rates',cfg,'rates.xlsx','sha')
    assert result['errors']==[] and result['records'][0]['rate']=='12.1250'
