"""Bounded spreadsheet adapters. No financial calculation or approval authority."""
import base64
import csv
import io
import json
import re
import subprocess
import sys
from datetime import date, datetime
from pathlib import PurePosixPath
from zipfile import ZipFile, BadZipFile
from .engine import PARSERS

MAX_BYTES=2*1024*1024
MAX_ROWS=2500
MAX_COLUMNS=50
MAX_SHEETS=8
MAX_XML=16*1024*1024
MAX_TEXT=2*1024*1024

# Aliases only suggest a mapping. A person must confirm both columns and values.
FIELDS={
 "rates":{
  "id":("Rate version",True,["rate id","rate version","rule id"]),
  "contract_id":("Contract reference",True,["contract","contract id","agreement id"]),
  "facility_id":("Facility reference",True,["facility","facility id","client","client id","customer"]),
  "role":("Role",True,["role","job","job title","discipline","specialty"]),
  "component":("Hourly component",True,["component","pay type","bill type","earning code","rate type"]),
  "rate":("Contract bill rate per hour",True,["bill rate","hourly rate","client rate","contract rate"]),
  "effective_from":("Effective from",True,["effective from","effective date","start date"]),
  "effective_to":("Effective end (optional)",False,["effective to","end date","expiration date"]),
  "authorization_date":("Source authorization date",True,["authorization date","authorized on","agreement approved","approved date"]),
  "supersedes":("Previous rate version (amendment)",False,["supersedes","previous rate","previous rate id"]),
  "currency":("Currency (optional; USD otherwise)",False,["currency","currency code"]),
 },
 "work":{
  "id":("Work record reference",True,["shift id","work id","timesheet id","record id"]),
  "facility_id":("Facility reference",True,["facility","facility id","client","client id","customer"]),
  "worker_id":("Worker reference (ID only)",True,["employee id","worker id","clinician id","resource id","employee","worker","clinician","resource"]),
  "assignment_id":("Assignment reference (optional)",False,["assignment id","placement id"]),
  "role":("Role",True,["role","job","job title","discipline","specialty"]),
  "component":("Hourly component",True,["component","pay type","bill type","earning code"]),
  "service_date":("Service date",True,["service date","shift date","work date","date worked"]),
  "service_end_date":("Service end date (overnight)",False,["service end date","shift end date"]),
  "hours":("Worked time",True,["hours","approved hours","worked hours","bill hours","minutes"]),
  "approval_state":("Time approval status",True,["approval","approval status","time status","approved"]),
  "approved_at":("Time approved date",False,["approved at","approved date","time approved date"]),
  "source_revision":("Source version (optional)",False,["revision","source revision","version"]),
 },
 "invoices":{
  "id":("Invoice line reference",True,["line id","invoice line id","billing line id","record id"]),
  "invoice_id":("Invoice or statement reference",True,["invoice #","invoice number","billing id","invoice id","statement id"]),
  "facility_id":("Facility reference",True,["facility","facility id","client","client id","customer"]),
  "shift_ids":("Work record references (optional)",False,["shift id","shift ids","work id","timesheet id"]),
  "worker_id":("Worker reference for matching",False,["employee id","worker id","clinician id","resource id"]),
  "service_date":("Service date for matching",False,["service date","shift date","work date"]),
  "role":("Role for matching",False,["role","job","discipline"]),
  "component":("Hourly component for matching",False,["component","pay type","bill type","earning code"]),
  "amount":("Signed invoice amount",True,["amount","line amount","net amount","billed amount","total"]),
  "hours":("Billed time",True,["hours","billed hours","bill hours","minutes"]),
  "unit_rate":("Invoice bill rate per hour",True,["bill rate","hourly rate","client rate","unit rate"]),
  "state":("Posting status",True,["status","posting status","invoice status"]),
  "document_kind":("Document type",True,["document type","type","document kind"]),
  "chain_id":("Credit/rebill chain reference",False,["chain id","credit chain","rebill chain"]),
  "currency":("Currency (optional; USD otherwise)",False,["currency","currency code"]),
 },
}
VALUES={
 "component":{"standard":["regular","reg","standard","straight time"],"overtime":["overtime","ot"],"holiday":["holiday","hol"],"weekend":["weekend"],"on_call":["on call","oncall"],"callback":["callback","call back"]},
 "approval_state":{"approved":["approved","yes","true","1"],"pending":["pending","no","false","0","unapproved"],"rejected":["rejected","denied"]},
 "state":{"posted":["posted","final","paid"],"draft":["draft","pending"],"void":["void","voided","cancelled"]},
 "document_kind":{"invoice":["invoice","statement","reverse billing","hospital statement"],"credit":["credit","credit memo"],"debit":["debit","adjustment","rebill"]},
}
DENIED={"name","firstname","lastname","fullname","phone","telephone","birthdate","resourcename","contactname","patient","patientname","patientid","diagnosis","diagnosiscode","clinicalnotes","medicalrecord","mrn","ssn","socialsecuritynumber","dateofbirth","dob","employeename","workername","clinicianname","email","address","bankaccount"}

class TableError(ValueError):
    def __init__(self,code,message):
        self.code=code
        super().__init__(message)

def token(value):
    return re.sub(r"[^a-z0-9]","",str(value).casefold())

def suggest(headers,category):
    result={}
    for field,(_,_,aliases) in FIELDS[category].items():
        candidates=[h for h in headers if token(h) in {token(a) for a in [field,*aliases]}]
        if len(candidates)==1:result[field]=candidates[0]
    return result

def value_suggestions(rows,headers,mapping):
    result={}
    for field in ["role",*VALUES]:
        if field not in mapping:continue
        index=headers.index(mapping[field])
        source_values=sorted({str(r["values"][index]).strip() for r in rows if r["values"][index]!=""})
        if len(source_values)>100:raise TableError("VALUE_LIMIT","Use at most 100 distinct values per mapped status, role or component.")
        result[field]={}
        for value in source_values:
            if field=="role":result[field][value]=value
            else:
                result[field][value]=next((k for k,aliases in VALUES[field].items() if token(value) in {token(k),*(token(a) for a in aliases)}),"")
    return result

def cell(value):
    if value is None:return ""
    if isinstance(value,(date,datetime)):return value.isoformat()
    value=str(value)
    if len(value)>2000 or any(ord(c)<32 and c not in "\t\n\r" for c in value):
        raise TableError("CELL_LIMIT","A cell is too long or contains unsupported control characters.")
    return value.strip()

def read_tables(data,filename):
    if not data or len(data)>MAX_BYTES:raise TableError("FILE_LIMIT","Choose a nonempty file no larger than 2 MiB.")
    suffix=filename.rsplit(".",1)[-1].lower()
    if suffix=="csv":
        try:
            text=data.decode("utf-16" if data[:2] in [b'\xff\xfe',b'\xfe\xff'] else "utf-8-sig")
            dialect=csv.Sniffer().sniff(text[:65536],delimiters=",;\t|")
            rows=[]
            for row in csv.reader(io.StringIO(text),dialect):
                if len(rows)>=MAX_ROWS+20 or len(row)>MAX_COLUMNS:raise TableError("ROW_LIMIT","Use at most 2,500 data rows and 50 columns per worksheet.")
                rows.append([cell(v) for v in row])
            return {"CSV":rows}
        except (UnicodeError,csv.Error):raise TableError("CSV_FORMAT","Use a UTF-8 or BOM-marked UTF-16 CSV with a consistent comma, semicolon, tab or pipe delimiter.") from None
    if suffix!="xlsx":raise TableError("FILE_TYPE","Choose .csv or .xlsx. Macros and legacy .xls files are not accepted.")
    from defusedxml import ElementTree as ET
    from openpyxl import load_workbook
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries=archive.infolist()
            if len(entries)>150 or len({e.filename for e in entries})!=len(entries) or sum(e.file_size for e in entries)>MAX_XML:
                raise TableError("ARCHIVE_LIMIT","Workbook exceeds the safe archive limits.")
            for entry in entries:
                name=entry.filename.lower()
                if entry.flag_bits&1 or name.startswith('/') or '..' in PurePosixPath(name).parts or entry.file_size>MAX_XML or entry.file_size>max(1,entry.compress_size)*200:
                    raise TableError("ARCHIVE_LIMIT","Workbook exceeds the safe archive limits or is encrypted.")
                if any(part in name for part in ['vbaproject','externallinks/','embeddings/','connections.xml','activex/','macrosheets/']):
                    raise TableError("ACTIVE_CONTENT","Remove macros, embedded objects, connections and external links.")
                if name.endswith(('.xml','.rels')):
                    with archive.open(entry) as stream:
                        for _,element in ET.iterparse(stream,events=("end",)):
                            if element.tag.rsplit('}',1)[-1]=='f' or element.attrib.get('TargetMode')=='External':
                                raise TableError("ACTIVE_CONTENT","Formula cells and external references are not accepted. Export values only.")
                            if element.tag.rsplit('}',1)[-1]=='mergeCell':raise TableError("MERGED_CELLS","Unmerge cells in the exported worksheets before importing.")
                            element.clear()
        workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            if len(workbook.worksheets)>MAX_SHEETS:raise TableError("SHEET_LIMIT","Use at most eight worksheets per workbook.")
            tables={}
            total=0;characters=0
            for sheet in workbook.worksheets:
                if sheet.sheet_state!='visible':raise TableError("HIDDEN_SHEET","Remove hidden worksheets from the export before importing.")
                sheet.reset_dimensions()
                rows=[]
                for row in sheet.iter_rows():
                    if len(rows)>=MAX_ROWS+20 or len(row)>MAX_COLUMNS:raise TableError("ROW_LIMIT","Use at most 2,500 data rows and 50 columns per worksheet.")
                    values=[cell(c.value) for c in row]
                    characters+=sum(len(value) for value in values)
                    if characters>MAX_TEXT:raise TableError('TEXT_LIMIT','Workbook contains too much expanded text. Export fewer columns or a shorter period.')
                    rows.append(values)
                    total+=len(row)
                    if total>100000:raise TableError("CELL_LIMIT","Workbook exceeds 100,000 cells.")
                tables[sheet.title]=rows
            return tables
        finally:workbook.close()
    except TableError:raise
    except Exception:
        raise TableError("XLSX_FORMAT","Workbook could not be safely read. Save a new values-only .xlsx export.") from None

def table_view(rows,category,header_row=None):
    if not rows:raise TableError("EMPTY_FILE","The selected worksheet is empty.")
    if header_row is None:
        def score(row):
            return len(suggest(row,category))*100+len([v for v in row if v])
        header_row=max(range(min(20,len(rows))),key=lambda i:score(rows[i]))+1
    if not isinstance(header_row,int) or not 1<=header_row<=min(20,len(rows)):
        raise TableError("HEADER_ROW","Choose a header row between 1 and 20.")
    headers=list(rows[header_row-1])
    while headers and not headers[-1]:headers.pop()
    if not headers or any(not h for h in headers) or len({token(h) for h in headers})!=len(headers):
        raise TableError("HEADER_FIELDS","Every column needs a distinct, nonempty header. Select the correct header row.")
    if any(token(h) in DENIED or token(h).startswith(('patient','diagnosis','clinical','medicalrecord')) for h in headers):
        raise TableError("SENSITIVE_FIELDS","Remove names, patient, clinical, identity, contact and banking columns before uploading.")
    data=[]
    for number,row in enumerate(rows[header_row:],header_row+1):
        if not any(row):continue
        if len(row)>len(headers) and any(row[len(headers):]):raise TableError("COLUMN_COUNT","A row has values beyond the selected header. Correct the source file.")
        data.append({"row_number":number,"values":row[:len(headers)]+['']*max(0,len(headers)-len(row))})
    if (not data and category!='invoices') or len(data)>MAX_ROWS:raise TableError("ROW_LIMIT","Use 1 to 2,500 data rows per selected worksheet.")
    mapping=suggest(headers,category)
    return {"header_row":header_row,"headers":headers,"rows":data,"mapping":mapping,"value_maps":value_suggestions(data,headers,mapping)}

def inspect_file(data,filename,category,sheet=None,header_row=None):
    if category not in FIELDS:raise TableError("CATEGORY","Choose rates, work or invoices.")
    tables=read_tables(data,filename)
    if sheet is not None and sheet not in tables:raise TableError("WORKSHEET","Select an available worksheet.")
    if sheet is None:
        sheet=max(tables,key=lambda name:max((len(suggest(row,category)) for row in tables[name][:20]),default=0))
    return {"worksheets":list(tables),"worksheet":sheet,**table_view(tables[sheet],category,header_row)}

def inspect_isolated(data,filename,category,sheet=None,header_row=None):
    payload={"data":base64.b64encode(data).decode(),"filename":filename,"category":category,"sheet":sheet,"header_row":header_row}
    try:
        with PARSERS:
            result=subprocess.run([sys.executable,'-m','backend.table_cli'],input=json.dumps(payload),text=True,capture_output=True,timeout=10,check=False,env={'PATH':'/usr/bin:/bin','LANG':'C.UTF-8','PYTHONNOUSERSITE':'1','PYTHONDONTWRITEBYTECODE':'1'})
        response=json.loads(result.stdout)
        if not response.get('ok'):raise TableError(response.get('code','PARSER_LIMIT'),response.get('message','Use a smaller values-only file.'))
        return response['table']
    except (subprocess.TimeoutExpired,ValueError,OSError) as error:
        if isinstance(error,TableError):raise
        raise TableError("PARSER_LIMIT","Parsing exceeded its time/resource limit. Use a smaller values-only file.") from None
