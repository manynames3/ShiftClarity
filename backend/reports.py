"""Current reviewed-run reports. Formatting never recalculates financial findings."""
import io
import re
from threading import BoundedSemaphore
from datetime import datetime, timezone
from decimal import Decimal
from html import escape
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
import reportlab
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table as PDFTable, TableStyle, KeepTogether
from .config import ROOT
from .engine import digest

STATES={'needs_review':'Needs review','confirmed':'Confirmed discrepancy','dismissed':'Not a discrepancy','evidence_requested':'Evidence requested'}
DISCLAIMER='Fictional evaluation only. Potential differences are not recovered revenue, payment evidence or confirmed invoice adjustments. Unevaluated work has no financial estimate.'
REPORTS=BoundedSemaphore(2)
NAVY='09273B'; TEAL='007E87'; MUTED='546A7B'; PALE='EAF6F5'

def label(value):
    return str(value).replace('_',' ').lower().capitalize()

def money(value):
    return f"{'-' if Decimal(value)<0 else ''}${abs(Decimal(value)):,.2f}"

def key(finding):
    direction='over' if finding['discrepancy_amount'].startswith('-') else 'under'
    return '/'.join([finding['dataset_id'],finding['facility_id'],finding['rule_version_id'],finding['reason_code'],direction])

def model(state, lineage=None, include_notes=False, cases=None, case_events=None):
    run={**state['run'],'fingerprint':digest(state['run']['fingerprint'])}; findings=[]; history=[]
    for source in run['findings']:
        review=state.get('reviews',{}).get(key(source),{})
        # A mismatched review fingerprint cannot attest to the current evidence.
        status=review.get('state','needs_review') if review.get('fingerprint')==source['fingerprint'] else 'needs_review'
        findings.append({**source,'review_state':status,'review_label':STATES[status]})
        for event in review.get('history',[]):
            history.append({'finding_id':source['id'],**{k:event.get(k,'') for k in ['sequence','action','actor','timestamp','owner','fingerprint']},'fingerprint':digest(event.get('fingerprint','')),'reason':event.get('reason','') if include_notes else '(notes omitted)'})
    for event in case_events or []:
        history.append({'finding_id':event['case_id'],'sequence':event['sequence'],'action':event['action'],'actor':event['actor'],'timestamp':event['created_at'].isoformat(),'owner':event['data'].get('owner',''),'reason':event['data'].get('reason','') if include_notes else '(notes omitted)','fingerprint':''})
    records=run['records']
    evaluated=[r for r in records if r['category'] in ['evaluated_aligned','evaluated_flagged']]
    evaluated_billed=sum((Decimal(r['invoiced_amount']) for r in evaluated),Decimal(0))
    confirmed={direction:sum((abs(Decimal(f['discrepancy_amount'])) for f in findings if f['review_state']=='confirmed' and (Decimal(f['discrepancy_amount'])<0)==(direction=='over')),Decimal(0)) for direction in ['under','over']}
    return {'run':run,'findings':findings,'records':records,'rates':state['inputs']['rules'],
            'source_counts':{kind:len(state['inputs'][kind]) for kind in ['work','invoices','rules']},'lineage':[{**row,'content_hash':row.get('sha256',row.get('content_hash',''))} for row in (lineage or [])],'history':history,'tasks':run['tasks'],'notes_included':include_notes,'cases':cases or [],
            'generated_at':datetime.now(timezone.utc).isoformat(timespec='seconds'),
            'stale':state.get('stale',False),'evaluated_billed':str(evaluated_billed),
            'confirmed':confirmed,'reviews_completed':sum(f['review_state'] in ['confirmed','dismissed'] for f in findings)}

def overview(data):
    run=data['run']; coverage=run['coverage']; m=run['manifest']
    return [
        ('Report','Contract-to-invoice review'),('Data boundary','Fictional records only'),
        ('Run reference',run['id']),('Report generated (UTC)',data['generated_at']),
        ('Service from (inclusive)',m['service_from']),('Service to (exclusive)',m['service_to_exclusive']),
        ('Source current through',m['as_of']),('Billing cutoff',m['billing_cutoff']),
        ('Result freshness','STALE - previous completed run retained; rerun validated evidence' if data['stale'] else 'Completed run - original source evidence'),
        ('Source records accepted',sum(data['source_counts'].values())),('Records received',coverage['received_records']),('Records evaluated',coverage['evaluated_records']),
        ('Hours received',Decimal(coverage['received_minutes'])/60),('Hours evaluated',Decimal(coverage['evaluated_minutes'])/60),
        ('Expected on evaluated work (USD)',Decimal(run['totals']['expected'])),
        ('Net invoiced on evaluated work (USD)',Decimal(data['evaluated_billed'])),
        ('All posted invoice lines, including unchecked work (USD)',Decimal(run['totals']['invoiced'])),
        ('Potential underbilling before review exclusions (USD)',Decimal(run['totals']['underbilling'])),
        ('Potential overbilling before review exclusions (USD)',Decimal(run['totals']['overbilling'])),
        ('Reviewer-confirmed underbilling difference (USD)',data['confirmed']['under']),
        ('Reviewer-confirmed overbilling difference (USD)',data['confirmed']['over']),
        ('Findings',len(data['findings'])),('Reviews completed',data['reviews_completed']),
        *[(STATES[state],sum(f['review_state']==state for f in data['findings'])) for state in STATES],
        ('Coverage tasks',len(run['tasks'])),('Reviewer notes included','Yes - explicitly requested' if data['notes_included'] else 'No'),
        ('Worker identities','Omitted'),('Method','Expected = exact approved hourly rate x worked minutes / 60; compare signed invoice allocations; money rounded by the shared engine.'),
        ('Coverage declaration','Declared complete' if m['invoice_coverage_complete'] else 'Incomplete - missing invoices cannot be treated as underbilling'),
        ('Controls','Parsed counts/totals are not independent proof of completeness.'),
        ('Limit',DISCLAIMER),('Calculation engine',run['engine_version']),('Run evidence fingerprint',run['fingerprint'])]

def xlsx(data):
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,headers,rows,widths):
        ws=wb.create_sheet(name);ws.append(headers)
        for row in rows:
            ws.append(row)
            # Literal strings, even when starting with =,+,-,@ or a URL. No formulas or hyperlinks.
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
                if isinstance(cell.value,Decimal):
                    cell.number_format='#,##0.00;[Red](#,##0.00)'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor=NAVY);cell.alignment=Alignment(vertical='center',wrap_text=True)
        ws.row_dimensions[1].height=32;ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
        for i,width in enumerate(widths,1):ws.column_dimensions[ws.cell(1,i).column_letter].width=width
        if ws.max_row>1:
            table=Table(displayName=re.sub('[^A-Za-z]','',name)+'Data',ref=ws.dimensions)
            table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
        ws.print_title_rows='1:1';ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.oddFooter.center.text='ShiftClarity | Fictional evaluation only | Page &P'
        return ws
    sheet('Overview',['Measure','Value'],overview(data),[62,100])
    sheet('Findings',['Finding','Facility','Review state','Review basis','Affected work records','Rate version','Expected USD','Net invoiced USD','Difference USD','Collectibility'],
          [[f['id'],f['facility_id'],f['review_label'],label(f['reason_code']),len(f['shift_ids']),f['rule_version_id'],Decimal(f['expected_amount']),Decimal(f['invoiced_amount']),Decimal(f['discrepancy_amount']),'Not established'] for f in data['findings']],[22,24,28,42,42,32,20,20,20,22])
    sheet('Work calculations',['Work reference','Facility','Service date','Role','Hourly component','Minutes','Hours','Coverage','Review basis','Rate version','Rate USD per hour','Expected USD','Net invoiced USD','Difference USD','Invoice allocation count'],
          [[r['work']['id'],r['work']['facility_id'],r['work']['service_date'],r['work']['role'],r['work']['component'],r['work']['worked_minutes'],Decimal(r['work']['worked_minutes'])/60,label(r['category']),label(r['reason_code']),r.get('rule',{}).get('id',''),Decimal(r['rule']['rate']) if r.get('rule') else None,*[Decimal(r[k]) if r[k] is not None else None for k in ['expected_amount','invoiced_amount','discrepancy_amount']],len(r['allocations'])] for r in data['records']],[22,22,18,18,22,12,12,25,42,32,22,20,20,20,70])
    sheet('Invoice allocations',['Work reference','Invoice reference','Line reference','Signed allocated USD','Credit/rebill chain'],[[a['shift_id'],a['invoice_id'],a['line_id'],Decimal(a['amount']),a.get('chain_id') or ''] for a in data['run']['allocations']],[24,35,40,25,35])
    sheet('Rate terms',['Rate version','Contract reference','Facility','Role','Hourly component','Rate USD per hour','Effective from','Effective to (exclusive)','Source authorized','System approved at','Author account','Approver account','Evidence reference'],
          [[r['id'],r['contract_id'],r['facility_id'],r['role'],r['component'],Decimal(r['rate']),r['effective_from'],r.get('effective_to') or '',r.get('authorization_date') or (r.get('approved_at') or '')[:10],r.get('approved_at',''),r['author_id'],r.get('approver_id',''),r['source_excerpt_id']] for r in data['rates']],[32,28,22,18,22,22,20,24,20,34,28,28,40])
    sheet('Source lineage',['Source category','Record reference','Filename','Worksheet','Row','Original SHA256','Source file reference'],
          [[l.get(k,'') for k in ['kind','record_id','filename','worksheet','row','content_hash','file_id']] for l in data['lineage']],[22,30,38,30,10,72,45])
    sheet('Review history',['Finding','Sequence','Action','Account','Recorded at','Task owner','Reason','Evidence fingerprint'],
          [[h.get(k,'') for k in ['finding_id','sequence','action','actor','timestamp','owner','reason','fingerprint']] for h in data['history']],[22,12,26,28,34,28,85,80])
    sheet('Coverage gaps',['Task','Work references','Coverage','Reason','Financial estimate'],
          [[t['id'],reference,label(t['category']),label(t['reason_code']),'No estimate applied'] for t in data['tasks'] for reference in t['shift_ids']],[22,42,28,55,28])
    sheet('Follow-up',['Case reference','Facility','Review basis','Latest status','Owner','Correction / evidence reference','Verified run','Last evidence run','Updated at'],[[c['id'],c['facility_id'],label(c['reason_code']),label(c['status']),c['owner'] or '',c['correction_reference'] or '',c['verified_run_id'] or '',c['last_run_id'],c['updated_at'].isoformat()] for c in data['cases']],[45,24,42,32,32,42,45,45,34])
    for ws,column in [(wb['Rate terms'],6),(wb['Work calculations'],11)]:
        for row in range(2,ws.max_row+1):ws.cell(row,column).number_format='#,##0.0000'
    output=io.BytesIO();wb.save(output);return output.getvalue()

def pdf(data):
    for name,filename in [('SC-Sans','Vera.ttf'),('SC-Bold','VeraBd.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,str(Path(reportlab.__file__).parent/'fonts'/filename)))
    glyphs=pdfmetrics.getFont('SC-Sans').face.charToGlyph
    output=io.BytesIO();width=7*inch
    styles=getSampleStyleSheet()
    for name in ['Normal','BodyText']:
        styles[name].fontName='SC-Sans';styles[name].fontSize=9;styles[name].leading=12;styles[name].textColor=colors.HexColor('#'+NAVY)
    styles.add(ParagraphStyle(name='SCtitle',fontName='SC-Bold',fontSize=29,leading=34,textColor=colors.HexColor('#'+NAVY),spaceAfter=12))
    styles.add(ParagraphStyle(name='SCsection',fontName='SC-Bold',fontSize=15,leading=20,textColor=colors.HexColor('#'+NAVY),spaceBefore=16,spaceAfter=6,keepWithNext=True))
    styles.add(ParagraphStyle(name='SCsmall',parent=styles['Normal'],fontSize=8,leading=11,textColor=colors.HexColor('#'+MUTED)))
    def p(value,style='Normal'):
        # Treat every external string as text, including ReportLab paragraph markup.
        text=''.join(c if c in '\n\t' or ord(c) in glyphs else f'[U+{ord(c):04X}]' for c in str(value))
        return Paragraph(escape(text).replace('\n','<br/>'),styles[style])
    def section(title):return p(title,'SCsection')
    def table(headers,rows,widths):
        cells=[[p(value,'SCsmall') for value in headers]]+[[p(value) for value in row] for row in rows]
        result=PDFTable(cells,colWidths=widths,repeatRows=1,hAlign='LEFT')
        result.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#'+PALE)),('LINEBELOW',(0,0),(-1,0),.6,colors.HexColor('#'+TEAL)),('LINEBELOW',(0,1),(-1,-1),.3,colors.HexColor('#DCE6EC')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7)]))
        return result
    run=data['run'];c=run['coverage'];m=run['manifest']
    story=[p('Billing review','SCtitle'),p(f"Service period: {m['service_from']} to {m['service_to_exclusive']} (end exclusive)",'BodyText'),p('Fictional evaluation | USD | Prepared '+data['generated_at'],'SCsmall'),Spacer(1,18)]
    if data['stale']:story.extend([p('STALE RESULTS: Previous completed run retained. Rerun validated evidence before making a decision.','BodyText'),Spacer(1,12)])
    story.extend([table(['Potential underbilling','Potential overbilling','Work evaluated'],[[money(run['totals']['underbilling']),money(run['totals']['overbilling']),f"{c['evaluated_records']} of {c['received_records']} records"]],[width*.34,width*.34,width*.32]),Spacer(1,9),p(DISCLAIMER,'SCsmall'),
                  section('What this review establishes'),p(f"{sum(data['source_counts'].values())} accepted source records: {data['source_counts']['work']} work, {data['source_counts']['invoices']} invoice lines, {data['source_counts']['rules']} rate versions."),p(f"{len(data['findings'])} finding{'s' if len(data['findings'])!=1 else ''}. Reviews completed: {data['reviews_completed']}. Coverage tasks: {len(run['tasks'])}. All finding states are included; potential totals above include dismissed findings before review exclusions."),
                  p(f"Reviewer-confirmed differences: {money(data['confirmed']['under'])} underbilling; {money(data['confirmed']['over'])} overbilling. Confirmation records a review, not an invoice adjustment or collected funds."),
                  table(['Evaluated-work comparison','Amount USD'],[['Expected under approved hourly terms',money(run['totals']['expected'])],['Signed invoice allocations on evaluated work',money(data['evaluated_billed'])],['All posted invoice lines, including unchecked work',money(run['totals']['invoiced'])]],[width*.76,width*.24]),
                  p('Do not subtract the all-posted total from evaluated expected billing. Coverage scopes differ.','SCsmall')])
    if not data['findings']:story.extend([section('Findings and supporting calculations'),p('No calculated findings. Review the coverage gaps before interpreting this as agreement.')])
    records={r['work']['id']:r for r in data['records']}
    for index,finding in enumerate(data['findings']):
        summary=[p(f"{finding['id']} | {finding['facility_id']} | {finding['review_label']}",'BodyText'),p(label(finding['reason_code'])+' | Rate version '+finding['rule_version_id'],'SCsmall'),table(['Expected','Net invoiced','Signed difference'],[[money(finding['expected_amount']),money(finding['invoiced_amount']),money(finding['discrepancy_amount'])]],[width/3]*3)]
        if index==0:summary.insert(0,section('Findings and supporting calculations'))
        block=[*summary,Spacer(1,6)]
        rows=[]
        for identifier in finding['shift_ids']:
            r=records[identifier];w=r['work']
            rows.append([identifier,w['service_date'],w['worked_minutes'],f"{Decimal(w['worked_minutes'])/60:.2f}",money(r['expected_amount']),money(r['invoiced_amount']),money(r['discrepancy_amount'])])
        block.append(table(['Work reference','Service date','Minutes','Hours*','Expected','Invoiced','Difference'],rows,[width*.18,width*.17,width*.11,width*.09,width*.15,width*.15,width*.15]))
        story.extend([KeepTogether(block)] if len(rows)<=8 else [KeepTogether(summary),*block[len(summary):]])
        refs=['; '.join(a['invoice_id']+' / '+a['line_id']+': '+a['amount'] for a in records[identifier]['allocations']) for identifier in finding['shift_ids']]
        story.append(p('Invoice allocations: '+'; '.join(filter(None,refs)) if any(refs) else 'Invoice allocations: none. Missing-billing findings depend on the explicit complete-source declaration and cutoff.','SCsmall'))
        story.append(p('Calculation: exact approved hourly rate x minutes / 60, less signed posted invoice allocations. *Hours display two decimals; calculations use whole minutes. Collectibility is not established.','SCsmall'))
        story.append(Spacer(1,12))
    story.append(section('Coverage and evidence limits'))
    story.append(p(f"{c['received_minutes']/60:.2f} hours received; {c['evaluated_minutes']/60:.2f} hours evaluated. Invoice coverage: {'declared complete' if m['invoice_coverage_complete'] else 'incomplete'}. Source as of {m['invoice_source_as_of']}; billing cutoff {m['billing_cutoff']}. Parsed controls do not establish completeness."))
    if data['tasks']:story.append(table(['Work references','Coverage gap','Estimate'],[[reference,label(t['reason_code']),'None'] for t in data['tasks'] for reference in t['shift_ids']],[width*.30,width*.54,width*.16]))
    else:story.append(p('No coverage tasks were generated for this run.'))
    story.append(section('Applicable approved terms'))
    for rate in data['rates']:
        story.append(KeepTogether([p(f"{rate['id']} | {rate['facility_id']} | {rate['role']} / {rate['component']} | ${Decimal(rate['rate']):,.4f} per hour"),p(f"Effective {rate['effective_from']} to {rate.get('effective_to') or 'open'} (end exclusive). Source authorized: {rate.get('authorization_date') or (rate.get('approved_at') or '')[:10]}. Actual system approval: {rate.get('approved_at') or 'not supplied'} by {rate.get('approver_id') or 'not supplied'}. Evidence: {rate['source_excerpt_id']}.",'SCsmall'),Spacer(1,7)]))
    story.append(section('Source references'))
    if data['lineage']:
        groups={}
        for line in data['lineage']:groups.setdefault((line.get('filename',''),line.get('worksheet',''),line.get('content_hash','')),[]).append(line)
        rows=[]
        for (filename,worksheet,sha),lines in groups.items():
            for offset in range(0,len(lines),12):
                rows.append([filename+' / '+worksheet+('\nSHA256: '+sha if offset==0 else ' (continued)'),'; '.join(line.get('record_id','')+' (row '+str(line.get('row',''))+')' for line in lines[offset:offset+12])])
        story.append(table(['Original file / worksheet / SHA256','Record references and source rows'],rows,[width*.46,width*.54]))
    else:story.append(p('Canonical source selection. Inspect the rate and invoice references above; spreadsheet row lineage is not available for this run.'))
    if data['cases']:
        story.extend([section('Latest follow-up for this period'),p('Status as of report generation. The financial calculations above retain the selected run’s original evidence. Alignment does not establish payment.','SCsmall'),table(['Facility / case','Latest status / owner','Correction or evidence reference'],[[c['facility_id']+' / '+c['id'],label(c['status'])+' / '+(c['owner'] or 'Not assigned'),(c['correction_reference'] or 'None')+(' / Verified: '+c['verified_run_id'] if c['verified_run_id'] else '')] for c in data['cases']],[width*.30,width*.31,width*.39])])
    story.append(section('Review history'))
    story.append(p('Reviewer free-text notes are '+('included by explicit request.' if data['notes_included'] else 'omitted. Request them explicitly only when authorized. Worker identities are omitted.'),'SCsmall'))
    if data['history']:
        for event in data['history']:
            story.append(p(f"{event['finding_id']} | {event['action']} | {event['actor']} | {event['timestamp']}",'BodyText'))
            if data['notes_included']:story.append(p(event['reason']))
            story.append(Spacer(1,5))
    else:story.append(p('No reviewer decisions have been recorded.'))
    story.append(KeepTogether([section('Report reference'),p('Run: '+run['id'],'SCsmall'),p('Evidence fingerprint: '+run['fingerprint'],'SCsmall'),p('Engine: '+run['engine_version']+' | Generated: '+data['generated_at'],'SCsmall')]))
    logo=ROOT/'dist/brand/shiftclarity-report-logo.png'
    if not logo.exists():logo=ROOT/'public/brand/shiftclarity-report-logo.png'
    def frame(canvas,doc):
        canvas.saveState()
        if logo.exists():canvas.drawImage(str(logo),.75*inch,10.28*inch,width=2*inch,height=.34*inch,mask='auto',preserveAspectRatio=True)
        else:
            canvas.setFont('SC-Bold',16);canvas.setFillColor(colors.HexColor('#'+NAVY));canvas.drawString(.75*inch,10.35*inch,'ShiftClarity')
        canvas.setStrokeColor(colors.HexColor('#DCE6EC'));canvas.line(.75*inch,10.15*inch,7.75*inch,10.15*inch)
        canvas.setFont('SC-Sans',8);canvas.setFillColor(colors.HexColor('#'+MUTED));canvas.drawString(.75*inch,.45*inch,'Fictional evaluation | Evidence for human review | ShiftClarity')
        canvas.drawRightString(7.75*inch,.45*inch,f'Page {doc.page}');canvas.restoreState()
    doc=SimpleDocTemplate(output,pagesize=(8.5*inch,11*inch),rightMargin=.75*inch,leftMargin=.75*inch,topMargin=1.08*inch,bottomMargin=.8*inch,title='ShiftClarity - Billing review',author='ShiftClarity',pageCompression=1)
    doc.build(story,onFirstPage=frame,onLaterPages=frame)
    return output.getvalue()

def mount_reports(app,settings):
    from fastapi import Request,HTTPException
    from fastapi.responses import Response
    from . import auth
    from .db import transaction,head,audit,review_workspace
    @app.get('/api/v1/runs/{run_id}/report.{format}')
    def export_report(run_id:str,format:str,request:Request,include_notes:bool=False):
        member=auth.principal(settings,request)
        if format not in ['xlsx','pdf']:raise HTTPException(404,'Report format not found.')
        if not REPORTS.acquire(blocking=False):raise HTTPException(429,'Two reports are being prepared. Try again shortly.')
        try:
            with transaction(settings,member['org_id']) as db:
                current=review_workspace(db,run_id)['state']
                row=db.execute('SELECT import_id FROM completed_runs WHERE id=%s',(run_id,)).fetchone()
                source=db.execute('SELECT data FROM table_sessions WHERE import_id=%s AND status=\'accepted\' ORDER BY created_at LIMIT 1',(row['import_id'],)).fetchone() if row else None
                from .workflow import case_rows
                cases=case_rows(db,run_id)
                events=db.execute('SELECT * FROM case_events WHERE case_id=ANY(%s) ORDER BY sequence',([c['id'] for c in cases],)).fetchall() if cases else []
                data=model(current,source['data'].get('lineage',[]) if source else [],include_notes,cases,events)
                audit(db,member['org_id'],member['actor'],'report_download',run_id,{'format':format,'notes_included':include_notes})
            content=xlsx(data) if format=='xlsx' else pdf(data)
            mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if format=='xlsx' else 'application/pdf'
            return Response(content,media_type=mime,headers={'Content-Disposition':f'attachment; filename="ShiftClarity-billing-review.{format}"'})
        finally:REPORTS.release()
